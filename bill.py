import os
import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime
from openpyxl import Workbook, load_workbook

FILE = "BillingData.xlsx"
HEADERS = ["Invoice", "Name", "Mobile", "Email", "Address",
           "Product", "Qty", "Price", "Total", "Date"]


def init_file():
    if not os.path.exists(FILE):
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)
        wb.save(FILE)

def safe(v):
    """None -> empty string, everything else -> string."""
    return "" if v is None else str(v)

def norm(v):
    """Normalize numbers like 9876543210.0 -> 9876543210."""
    s = safe(v).strip()
    return s[:-2] if s.endswith(".0") else s

def get_rows():
    """Read all bills. Returns list of rows, or None on error."""
    try:
        wb = load_workbook(FILE)
        ws = wb.active
        return [r for r in ws.iter_rows(min_row=2, values_only=True)
                if any(c is not None for c in r)]
    except PermissionError:
        messagebox.showerror("Error", f"Please close {FILE} in Excel and try again.")
    except Exception as e:
        messagebox.showerror("Error", f"Could not read file:\n{e}")
    return None

def show_in_table(rows):
    for item in tree.get_children():
        tree.delete(item)
    for row in rows:
        tree.insert("", "end", values=[safe(c) for c in row])

def fill_form(r):
    r = list(r) + [None] * (10 - len(r))
    invoice_var.set(safe(r[0]))
    name_var.set(safe(r[1]))
    mobile_var.set(norm(r[2]))
    email_var.set(safe(r[3]))
    txt_address.delete("1.0", "end")
    txt_address.insert("1.0", safe(r[4]))
    product_var.set(safe(r[5]))
    qty_var.set(safe(r[6]))
    price_var.set(safe(r[7]))
    total_var.set(safe(r[8]))


init_file()

root = tk.Tk()
root.title("Smart Billing System ")
root.geometry("1300x800")
root.configure(bg="#dfeffc")

style = ttk.Style()
style.theme_use("clam")


invoice_var = tk.StringVar()
name_var = tk.StringVar()
mobile_var = tk.StringVar()
email_var = tk.StringVar()
product_var = tk.StringVar()
qty_var = tk.StringVar()
price_var = tk.StringVar()
total_var = tk.StringVar()


def generate_invoice():
    invoice_var.set("INV" + datetime.now().strftime("%Y%m%d%H%M%S"))

def calculate_total(show_error=True):
    try:
        total = float(qty_var.get()) * float(price_var.get())
        total_var.set(f"{total:.2f}")
        return True
    except ValueError:
        if show_error:
            messagebox.showerror("Error", "Enter valid quantity and price")
        return False

def save_bill():
    if not name_var.get().strip():
        messagebox.showerror("Error", "Enter customer name")
        return
    if not mobile_var.get().strip():
        messagebox.showerror("Error", "Enter mobile number")
        return
    if not calculate_total():
        return
    if not invoice_var.get():
        generate_invoice()

    try:
        wb = load_workbook(FILE)
        ws = wb.active
        ws.append([
            invoice_var.get(),
            name_var.get().strip(),
            norm(mobile_var.get()),
            email_var.get().strip(),
            txt_address.get("1.0", "end").strip(),
            product_var.get().strip(),
            qty_var.get(),
            price_var.get(),
            total_var.get(),
            datetime.now().strftime("%d-%m-%Y"),
        ])
        wb.save(FILE)
    except PermissionError:
        messagebox.showerror("Error", f"Close {FILE} in Excel, then save again.")
        return
    except Exception as e:
        messagebox.showerror("Error", f"Could not save:\n{e}")
        return

    load_bills()
    messagebox.showinfo("Success", "Bill Saved")

def load_bills():
    """View Bills: show every saved bill."""
    rows = get_rows()
    if rows is not None:
        show_in_table(rows)

search_win = None

def search_bill():
    """Open a separate window to search purchases by mobile number."""
    global search_win
    if search_win is not None and search_win.winfo_exists():
        search_win.lift()
        search_win.focus_force()
        return

    win = tk.Toplevel(root)
    search_win = win
    win.title("Search Bill - Customer Purchases")
    win.geometry("900x520")
    win.configure(bg="#dfeffc")

    tk.Label(win, text="SEARCH CUSTOMER PURCHASES", bg="#0b4ea2", fg="white",
             font=("Arial", 18, "bold")).pack(fill="x", ipady=10)

    top = tk.Frame(win, bg="#dfeffc")
    top.pack(fill="x", padx=20, pady=15)

    tk.Label(top, text="Mobile Number:", bg="#dfeffc",
             font=("Arial", 12, "bold")).pack(side="left")
    s_mobile = tk.StringVar(value=norm(mobile_var.get()))
    entry = tk.Entry(top, textvariable=s_mobile, width=25, font=("Arial", 12))
    entry.pack(side="left", padx=10)
    entry.focus_set()

    info = tk.Label(win, text="", bg="#dfeffc", font=("Arial", 11, "bold"),
                    justify="left", anchor="w")
    info.pack(fill="x", padx=20)

    cols = ("Invoice", "Date", "Product", "Qty", "Price", "Total")
    s_tree = ttk.Treeview(win, columns=cols, show="headings")
    widths = (170, 100, 250, 70, 100, 110)
    for c, w in zip(cols, widths):
        s_tree.heading(c, text=c)
        s_tree.column(c, width=w)
    s_tree.pack(fill="both", expand=True, padx=20, pady=10)

    grand = tk.Label(win, text="", bg="#dfeffc", fg="#0b4ea2",
                     font=("Arial", 13, "bold"))
    grand.pack(pady=(0, 10))

    def do_search(event=None):
        for item in s_tree.get_children():
            s_tree.delete(item)
        info.config(text="")
        grand.config(text="")

        mobile = norm(s_mobile.get())
        if not mobile:
            messagebox.showwarning("Search", "Enter a mobile number", parent=win)
            return

        rows = get_rows()
        if rows is None:
            return

        matches = [r for r in rows if norm(r[2]) == mobile]
        if not matches:
            messagebox.showerror("Not Found",
                                 "No bill found for this mobile number", parent=win)
            return

        last = list(matches[-1]) + [None] * 10
        info.config(text=f"Name: {safe(last[1])}    Email: {safe(last[3])}\n"
                         f"Address: {safe(last[4])}")

        grand_total = 0.0
        for r in matches:
            r = list(r) + [None] * 10
            s_tree.insert("", "end", values=(
                safe(r[0]), safe(r[9]), safe(r[5]),
                safe(r[6]), safe(r[7]), safe(r[8])))
            try:
                grand_total += float(r[8])
            except (TypeError, ValueError):
                pass

        grand.config(text=f"Total Bills: {len(matches)}     "
                          f"Total Spent: {grand_total:.2f}")

    ttk.Button(top, text="🔍 Search", command=do_search).pack(side="left", padx=5)
    ttk.Button(top, text="Close", command=win.destroy).pack(side="left", padx=5)
    entry.bind("<Return>", do_search)

    if s_mobile.get():
        do_search()

def on_row_select(event):
    sel = tree.selection()
    if sel:
        fill_form(tree.item(sel[0], "values"))

def clear_form():
    for v in (invoice_var, name_var, mobile_var, email_var,
              product_var, qty_var, price_var, total_var):
        v.set("")
    txt_address.delete("1.0", "end")

def send_email():
    messagebox.showinfo("Email", "Add your Gmail SMTP details here to send invoices.")


header = tk.Frame(root, bg="#0b4ea2", height=80)
header.pack(fill="x")

tk.Label(header, text="SMART BILLING SYSTEM ", bg="#0b4ea2", fg="white",
         font=("Arial", 24, "bold")).pack(side="left", padx=20)

clock = tk.Label(header, bg="#0b4ea2", fg="white", font=("Arial", 12, "bold"))
clock.pack(side="right", padx=20)

def update_clock():
    clock.config(text=datetime.now().strftime("%d-%m-%Y %H:%M:%S"))
    root.after(1000, update_clock)

update_clock()


cust = tk.LabelFrame(root, text="Customer Information",
                     bg="#e8f5e9", font=("Arial", 12, "bold"))
cust.place(x=20, y=100, width=580, height=280)

bill = tk.LabelFrame(root, text="Billing Information",
                     bg="#fff3e0", font=("Arial", 12, "bold"))
bill.place(x=650, y=100, width=620, height=280)


tk.Label(cust, text="Name", bg="#e8f5e9").grid(row=0, column=0, padx=10, pady=10)
tk.Entry(cust, textvariable=name_var, width=40).grid(row=0, column=1)

tk.Label(cust, text="Mobile", bg="#e8f5e9").grid(row=1, column=0, padx=10, pady=10)
tk.Entry(cust, textvariable=mobile_var, width=40).grid(row=1, column=1)

tk.Label(cust, text="Email", bg="#e8f5e9").grid(row=2, column=0, padx=10, pady=10)
tk.Entry(cust, textvariable=email_var, width=40).grid(row=2, column=1)

tk.Label(cust, text="Address", bg="#e8f5e9").grid(row=3, column=0)
txt_address = tk.Text(cust, width=35, height=5)
txt_address.grid(row=3, column=1)


labels = [
    ("Invoice No", invoice_var),
    ("Product", product_var),
    ("Quantity", qty_var),
    ("Unit Price", price_var),
    ("Total Amount", total_var),
]
for i, (txt, var) in enumerate(labels):
    tk.Label(bill, text=txt, bg="#fff3e0").grid(row=i, column=0, padx=10, pady=10)
    tk.Entry(bill, textvariable=var, width=40).grid(row=i, column=1)

ttk.Button(bill, text="Generate Invoice",
           command=generate_invoice).grid(row=0, column=2, padx=10)
ttk.Button(bill, text="Calculate Total",
           command=calculate_total).grid(row=4, column=2, padx=10)

btnframe = tk.Frame(root, bg="#dfeffc")
btnframe.place(x=20, y=410, width=1260, height=90)

buttons = [
    ("💾 Save Bill", save_bill),
    ("🔍 Search Bill", search_bill),
    ("📋 View Bills", load_bills),
    ("📧 Send Email", send_email),
    ("🧹 Clear Form", clear_form),
    ("❌ Exit", root.destroy),
]
for i, (txt, cmd) in enumerate(buttons):
    ttk.Button(btnframe, text=txt, command=cmd).grid(row=0, column=i, padx=10, pady=20)

tree = ttk.Treeview(root, columns=HEADERS, show="headings")
for c in HEADERS:
    tree.heading(c, text=c)
    tree.column(c, width=120)

scroll = ttk.Scrollbar(root, orient="vertical", command=tree.yview)
tree.configure(yscrollcommand=scroll.set)
tree.place(x=20, y=530, width=1240, height=230)
scroll.place(x=1260, y=530, width=20, height=230)
tree.bind("<<TreeviewSelect>>", on_row_select)

load_bills()
root.mainloop()